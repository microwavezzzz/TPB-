import os
import re
import json
import openpyxl

SCHEDULE_EXCEL_PATH = r"D:\unduhhh\(ADMIN ONLY) JADWAL & INFORMASI PERKULIAHAN-PRAKTIKUM-CORE PRODI TPB GASAL 2026-2027.xlsx"
STUDENT_EXCEL_PATH_1 = r"D:\unduhhh\[Share Mahasiswa] Pembagian Kelas TPB, PIK, AAPP, dan Core Prodi (1).xlsx"
STUDENT_EXCEL_PATH_2 = r"D:\unduhhh\[Share Mahasiswa] Pembagian Kelas TPB, PIK, AAPP, dan Core Prodi.xlsx"

def normalize_day(day_str):
    if not day_str:
        return ""
    d = day_str.strip().lower()
    if "senin" in d:
        return "Senin"
    if "selasa" in d:
        return "Selasa"
    if "rabu" in d:
        return "Rabu"
    if "kamis" in d:
        return "Kamis"
    if "jum" in d or "jumat" in d:
        return "Jumat"
    if "sabtu" in d:
        return "Sabtu"
    if "minggu" in d:
        return "Minggu"
    return day_str.strip().capitalize()

def parse_schedule_time(schedule_str):
    if not schedule_str:
        return {"day": "", "start": "", "end": "", "raw": ""}
    
    raw = str(schedule_str).strip()
    day = ""
    start = ""
    end = ""
    
    parts = raw.split(",")
    if len(parts) >= 2:
        day = normalize_day(parts[0])
        time_part = ",".join(parts[1:]).strip()
    else:
        match_day = re.match(r"^(Senin|Selasa|Rabu|Kamis|Jum['’]?at|Sabtu|Minggu)\s*[,:\s]?\s*(.*)", raw, re.IGNORECASE)
        if match_day:
            day = normalize_day(match_day.group(1))
            time_part = match_day.group(2).strip()
        else:
            time_part = raw

    time_match = re.search(r"(\d{1,2}[.:]\d{2})\s*[-–—s/d]+\s*(\d{1,2}[.:]\d{2})", time_part)
    if time_match:
        s_h, s_m = time_match.group(1).replace(".", ":").split(":")
        e_h, e_m = time_match.group(2).replace(".", ":").split(":")
        start = f"{int(s_h):02d}:{s_m}"
        end = f"{int(e_h):02d}:{e_m}"
    
    return {
        "day": day,
        "start": start,
        "end": end,
        "raw": raw
    }

def extract_hyperlink(cell):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    val = str(cell.value or "")
    if "http" in val:
        match = re.search(r"(https?://[^\s]+)", val)
        if match:
            return match.group(1)
    return ""

def parse_students_data():
    path = STUDENT_EXCEL_PATH_1 if os.path.exists(STUDENT_EXCEL_PATH_1) else STUDENT_EXCEL_PATH_2
    if not os.path.exists(path):
        return []

    wb = openpyxl.load_workbook(path, data_only=True)
    if "DATA" not in wb.sheetnames:
        return []

    ws = wb["DATA"]
    students = []
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not row or (row[0] is None and row[1] is None and row[2] is None):
            continue
        
        nim_val = row[1]
        if nim_val is not None:
            try:
                nim = str(int(nim_val)).strip()
            except:
                nim = str(nim_val).strip()
        else:
            nim = ""

        name = str(row[2] or "").strip()
        prodi = str(row[3] or "").strip()
        tpb_class = str(row[4] or "").strip()
        if tpb_class.isdigit():
            tpb_class = f"TPB {int(tpb_class):02d}"
        elif re.match(r"^TPB\s*(\d+)$", tpb_class, re.IGNORECASE):
            num = int(re.search(r"\d+", tpb_class).group(0))
            tpb_class = f"TPB {num:02d}"

        fakultas = str(row[5] or "").strip()
        beasiswa = str(row[6] or "").strip()
        core1 = str(row[7] or "").strip()
        core2 = str(row[8] or "").strip()
        pik = str(row[9] or "").strip()
        aapp = str(row[10] or "").strip()

        if name or nim:
            students.append({
                "nim": nim,
                "name": name,
                "prodi": prodi,
                "tpb_class": tpb_class,
                "fakultas": fakultas,
                "beasiswa": beasiswa,
                "core1": core1,
                "core2": core2,
                "pik": pik,
                "aapp": aapp
            })

    return students

def parse_schedule_data(excel_path=SCHEDULE_EXCEL_PATH):
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Schedule Excel file not found at: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    all_schedules = []
    classes_set = set()
    prodis_set = set()
    lecturers_set = set()
    rooms_set = set()
    core_prodi_classes = set()

    for sheet_name in wb.sheetnames:
        if sheet_name == "HOME":
            continue
        
        ws = wb[sheet_name]
        rows = list(ws.iter_rows())
        if not rows:
            continue

        # ==========================================
        # SPECIAL HANDLING: CORE PRODI SHEET
        # ==========================================
        if sheet_name == "CORE PRODI":
            header_idx = 4
            for r_idx in range(header_idx + 1, len(rows)):
                row = rows[r_idx]
                if not any(cell.value is not None for cell in row):
                    continue
                
                vals = [cell.value for cell in row]
                # Col 0: KLUSTER TPB, Col 1: KELOMPOK, Col 2: PRODI, Col 3: FAKULTAS
                # Col 4: KODE MK, Col 5: MATA KULIAH, Col 6: SKS, Col 7: KELAS
                # Col 8: NAMA DOSEN, Col 10: Link, Col 12: Pilihan Jadwal, Col 13: Ruang
                prodi_val = str(vals[2]).strip() if len(vals) > 2 and vals[2] is not None else ""
                code_val = str(vals[4]).strip() if len(vals) > 4 and vals[4] is not None else ""
                course_val = str(vals[5]).strip() if len(vals) > 5 and vals[5] is not None else ""
                sks_val = str(vals[6]).strip() if len(vals) > 6 and vals[6] is not None else "2.0"
                class_core = str(vals[7]).strip() if len(vals) > 7 and vals[7] is not None else ""
                lecturer_val = str(vals[8]).strip() if len(vals) > 8 and vals[8] is not None else ""
                link_val = extract_hyperlink(row[10]) if len(row) > 10 else ""
                schedule_val = str(vals[12]).strip() if len(vals) > 12 and vals[12] is not None else ""
                room_val = str(vals[13]).strip() if len(vals) > 13 and vals[13] is not None else ""

                if not course_val and not prodi_val:
                    continue

                parsed_time = parse_schedule_time(schedule_val)
                item_id = f"CORE_{r_idx}_{prodi_val}_{class_core}"

                item = {
                    "id": item_id,
                    "sheet": "CORE PRODI",
                    "category": "Core Prodi",
                    "class_name": f"Core {prodi_val} - Kelas {class_core}",
                    "core_prodi_name": prodi_val,
                    "core_class": class_core,
                    "course_name": course_val or "Core Prodi",
                    "course_code": code_val,
                    "sks": sks_val,
                    "lecturer": lecturer_val if lecturer_val != "None" else "",
                    "day": parsed_time["day"],
                    "start_time": parsed_time["start"],
                    "end_time": parsed_time["end"],
                    "raw_schedule": schedule_val,
                    "room": room_val if room_val != "None" else "",
                    "prodi": prodi_val,
                    "link": link_val,
                    "contact": ""
                }

                all_schedules.append(item)
                if prodi_val:
                    prodis_set.add(prodi_val)
                    core_prodi_classes.add(f"Core {prodi_val} - Kelas {class_core}")
                if room_val and room_val != "None":
                    rooms_set.add(room_val)
            continue

        # ==========================================
        # SPECIAL HANDLING: AAPP & PIK SHEETS
        # ==========================================
        if sheet_name in ["AAPP", "PIK"]:
            header_idx = 3
            for r_idx in range(header_idx + 1, len(rows)):
                row = rows[r_idx]
                if not any(cell.value is not None for cell in row):
                    continue
                vals = [cell.value for cell in row]
                course_val = str(vals[0]).strip() if len(vals) > 0 and vals[0] is not None else sheet_name
                class_val = str(vals[1]).strip() if len(vals) > 1 and vals[1] is not None else ""
                prodi_val = str(vals[2]).strip() if len(vals) > 2 and vals[2] is not None else ""
                code_val = str(vals[3]).strip() if len(vals) > 3 and vals[3] is not None else ""
                lecturer_val = str(vals[4]).strip() if len(vals) > 4 and vals[4] is not None else ""
                sks_val = str(vals[5]).strip() if len(vals) > 5 and vals[5] is not None else "2.0"
                link_val = extract_hyperlink(row[6]) if len(row) > 6 else ""
                schedule_val = str(vals[7]).strip() if len(vals) > 7 and vals[7] is not None else ""
                room_val = str(vals[8]).strip() if len(vals) > 8 and vals[8] is not None else ""

                if not class_val:
                    continue

                parsed_time = parse_schedule_time(schedule_val)
                item_id = f"{sheet_name}_{r_idx}_{class_val}"

                item = {
                    "id": item_id,
                    "sheet": sheet_name,
                    "category": "Core Prodi" if sheet_name == "AAPP" else "MKWU",
                    "class_name": class_val,
                    "course_name": course_val,
                    "course_code": code_val,
                    "sks": sks_val,
                    "lecturer": lecturer_val if lecturer_val != "None" else "",
                    "day": parsed_time["day"],
                    "start_time": parsed_time["start"],
                    "end_time": parsed_time["end"],
                    "raw_schedule": schedule_val,
                    "room": room_val if room_val != "None" else "",
                    "prodi": prodi_val,
                    "link": link_val,
                    "contact": ""
                }

                all_schedules.append(item)
                classes_set.add(class_val)
                if room_val and room_val != "None":
                    rooms_set.add(room_val)
            continue

        # ==========================================
        # STANDARD TPB SHEETS (MATEMATIKA, FISIKA, KIMIA, TBS, PRAKTIKUM, DTD, PK, DLL)
        # ==========================================
        header_idx = -1
        header_map = {}
        for r_idx, row in enumerate(rows[:8]):
            row_vals = [str(cell.value).strip().upper() if cell.value is not None else "" for cell in row]
            if any("KELAS" in v or "MATA KULIAH" in v or "KLUSTER" in v or "NAMA MK" in v for v in row_vals):
                header_idx = r_idx
                for c_idx, val in enumerate(row_vals):
                    if not val:
                        continue
                    if "KELAS" in val or "KELOMPOK" in val:
                        if "header_class" not in header_map:
                            header_map["header_class"] = c_idx
                    if "PRODI" in val:
                        if "header_prodi" not in header_map:
                            header_map["header_prodi"] = c_idx
                    if "MATA KULIAH" in val or "NAMA MK" in val:
                        if "header_course" not in header_map:
                            header_map["header_course"] = c_idx
                    if "KODE" in val:
                        if "header_code" not in header_map:
                            header_code = c_idx
                    if "SKS" in val:
                        if "header_sks" not in header_map:
                            header_map["header_sks"] = c_idx
                    if "DOSEN" in val:
                        if "header_lecturer" not in header_map:
                            header_map["header_lecturer"] = c_idx
                    if "JADWAL" in val or "WAKTU" in val:
                        if "header_schedule" not in header_map:
                            header_map["header_schedule"] = c_idx
                    if "RUANG" in val or "LABORATORIUM" in val:
                        if "header_room" not in header_map:
                            header_map["header_room"] = c_idx
                    if "LINK" in val or "WA" in val or "GCR" in val or "GOOGLECLASSROOM" in val:
                        if "header_link" not in header_map:
                            header_map["header_link"] = c_idx
                    if "KONTAK" in val or "ASPRAK" in val or "EMAIL" in val:
                        if "header_contact" not in header_map:
                            header_map["header_contact"] = c_idx
                break
        
        if header_idx == -1:
            continue

        category = "Kuliah"
        if "PRAK" in sheet_name.upper():
            category = "Praktikum"
        elif sheet_name.upper() in ["POLA SEHAT", "B. INDO", "B. INGGRIS"]:
            category = "MKWU"

        last_prodi = ""
        for r_idx in range(header_idx + 1, len(rows)):
            row = rows[r_idx]
            if not any(cell.value is not None for cell in row):
                continue
            
            def get_val(col_key, default=""):
                if col_key in header_map and header_map[col_key] < len(row):
                    v = row[header_map[col_key]].value
                    return str(v).strip() if v is not None else default
                return default

            def get_link(col_key):
                if col_key in header_map and header_map[col_key] < len(row):
                    cell = row[header_map[col_key]]
                    return extract_hyperlink(cell)
                return ""

            class_val = get_val("header_class")
            course_val = get_val("header_course")
            prodi_val = get_val("header_prodi")

            if not class_val and not course_val:
                continue

            if prodi_val:
                last_prodi = prodi_val
            else:
                prodi_val = last_prodi

            code_val = get_val("header_code")
            sks_val = get_val("header_sks")
            lecturer_val = get_val("header_lecturer")
            schedule_val = get_val("header_schedule")
            room_val = get_val("header_room")
            link_val = get_link("header_link") or get_val("header_link")
            contact_val = get_val("header_contact")

            if not schedule_val:
                for c in row:
                    if c.value and any(d in str(c.value).lower() for d in ["senin", "selasa", "rabu", "kamis", "jum", "sabtu"]):
                        schedule_val = str(c.value).strip()
                        break
            
            if not room_val:
                for c in row:
                    if c.value and any(k in str(c.value).upper() for k in ["GK1", "LAB", "LABTEK", "GKU", "R."]):
                        room_val = str(c.value).strip()
                        break

            parsed_time = parse_schedule_time(schedule_val)

            norm_class = class_val.strip()
            if norm_class.isdigit():
                norm_class = f"TPB {int(norm_class):02d}"
            elif re.match(r"^TPB\s*(\d+)$", norm_class, re.IGNORECASE):
                num = int(re.search(r"\d+", norm_class).group(0))
                norm_class = f"TPB {num:02d}"

            if not course_val:
                course_val = sheet_name

            item_id = f"{sheet_name}_{r_idx}_{norm_class}"

            item = {
                "id": item_id,
                "sheet": sheet_name,
                "category": category,
                "class_name": norm_class,
                "course_name": course_val,
                "course_code": code_val,
                "sks": sks_val,
                "lecturer": lecturer_val if lecturer_val != "None" else "",
                "day": parsed_time["day"],
                "start_time": parsed_time["start"],
                "end_time": parsed_time["end"],
                "raw_schedule": schedule_val,
                "room": room_val if room_val != "None" else "",
                "prodi": prodi_val,
                "link": link_val if "http" in link_val else "",
                "contact": contact_val
            }

            all_schedules.append(item)
            if norm_class:
                classes_set.add(norm_class)
            if prodi_val:
                for p in prodi_val.split("\n"):
                    p_clean = p.strip()
                    if p_clean:
                        prodis_set.add(p_clean)
            if lecturer_val and lecturer_val != "None":
                for l in re.split(r"[\n,;]|\d+\.\s*", lecturer_val):
                    l_clean = l.strip()
                    if l_clean and len(l_clean) > 3 and "PJ" not in l_clean:
                        lecturers_set.add(l_clean)
            if room_val and room_val != "None":
                rooms_set.add(room_val)

    def class_sort_key(c):
        m = re.match(r"TPB\s*(\d+)", c, re.IGNORECASE)
        if m:
            return (0, int(m.group(1)))
        return (1, c)

    sorted_classes = sorted(list(classes_set), key=class_sort_key)
    sorted_rooms = sorted(list(rooms_set))
    sorted_lecturers = sorted(list(lecturers_set))
    sorted_prodis = sorted(list(prodis_set))
    sorted_core_classes = sorted(list(core_prodi_classes))

    students_list = parse_students_data()

    result = {
        "metadata": {
            "source_schedule": os.path.basename(excel_path),
            "total_records": len(all_schedules),
            "total_classes": len(sorted_classes),
            "total_core_classes": len(sorted_core_classes),
            "total_rooms": len(sorted_rooms),
            "total_lecturers": len(sorted_lecturers),
            "total_students": len(students_list)
        },
        "classes": sorted_classes,
        "core_classes": sorted_core_classes,
        "rooms": sorted_rooms,
        "lecturers": sorted_lecturers,
        "prodis": sorted_prodis,
        "schedules": all_schedules,
        "students": students_list
    }

    return result

if __name__ == "__main__":
    data = parse_schedule_data()
    print("Parsed records:", data["metadata"]["total_records"])
    print("TPB Classes count:", len(data["classes"]))
    print("Core Classes count:", len(data["core_classes"]))
    print("Students count:", data["metadata"]["total_students"])
    
    os.makedirs("d:/TPB/data", exist_ok=True)
    with open("d:/TPB/data/schedule_data.json", "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print("Saved to d:/TPB/data/schedule_data.json")
